"""
Excel export for LPSE Monitor.

Exports relevant packages into ONE reused Excel workbook that senior staff
already work with, rather than generating a new file every time. The two
datasets the rest of the app keeps separate (see PROJECT_STATUS.md) export
separately here too: `DATASET_LELANG` ("Daftar Lengkap", the full /lelang
list) and `DATASET_BERANDA` ("Ringkasan Beranda", the homepage summary
with Akhir Pendaftaran) each have their own file/sheet/column/mode
configuration (`excel_config`, keyed by dataset) and their own
row-tracking (`excel_generated_rows`, keyed by file+sheet).

Safety rules this module enforces:
  1. ALWAYS back up the target file before touching it (see backup_file).
  2. NEVER write outside the configured sheet/start-cell area - other
     sheets, and any cell in this sheet the app didn't itself write, are
     left completely alone.
  3. Only ever clear/overwrite rows THIS APP previously wrote, tracked in
     the `excel_generated_rows` table by (file_path, sheet_name,
     row_number) - never "row N happens to look empty, so I'll use it."
  4. Use openpyxl (not pandas) to load/save, which preserves existing
     formatting, formulas, and every other sheet untouched.

Two write modes (user picks per export, from the Excel tab):
  - "replace":     clear every row this app previously wrote in that sheet,
                    then write the full current relevant-package list fresh.
  - "append_new":  leave existing rows alone, only add rows for packages
                    that don't have a tracked row yet.

Excel Table support (opt-in via `use_excel_table`, on by default): rather
than leaving plain values sitting in cells, the export also writes/updates
a native Excel Table (banded rows, filter dropdowns already turned on) -
so opening the workbook shows a ready-made table, and re-exporting later
updates that SAME table's range instead of leaving a stale one behind or
creating a second one. If the header row is missing text in any exported
column, or a DIFFERENT table already covers the same cells (e.g. someone
made their own Table over this range by hand), table-ification is skipped
for that export and only plain values are written - see _apply_excel_table.

Column mapping is configurable (which columns, in a fixed sensible order -
see DEFAULT_COLUMNS) rather than hardcoded. One column worth flagging:
"Pagu Anggaran" isn't included, because it isn't available from either
scraper (see PROJECT_STATUS.md limitations) - left out rather than filled
in with a misleading duplicate of HPS.
"""

from __future__ import annotations

import shutil
import sqlite3
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Optional

import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import coordinate_to_tuple, range_boundaries
from openpyxl.worksheet.table import Table, TableColumn, TableStyleInfo

from database import database as db

PROJECT_ROOT = Path(__file__).resolve().parent.parent
BACKUPS_DIR = PROJECT_ROOT / "backups"

DATASET_LELANG = "lelang"
DATASET_BERANDA = "beranda"

DATASET_LABELS = {
    DATASET_LELANG: "Daftar Lengkap",
    DATASET_BERANDA: "Ringkasan Beranda",
}
DEFAULT_SHEET_NAMES = {
    DATASET_LELANG: "Daftar Lengkap",
    DATASET_BERANDA: "Ringkasan Beranda",
}
# Excel table names can't contain spaces and must be unique per workbook.
TABLE_NAMES = {
    DATASET_LELANG: "TabelDaftarLengkap",
    DATASET_BERANDA: "TabelRingkasanBeranda",
}

COLUMN_LABELS = {
    "no": "No.",
    "nama_paket": "Nama Paket",
    "lpse": "LPSE",
    "jenis_pengadaan": "Jenis Pengadaan",
    "status": "Status",
    "bagian": "Bagian",
    "kategori": "Kategori",
    "hps": "HPS",
    "akhir_pendaftaran": "Akhir Pendaftaran",
    "tanggal_ditemukan": "Tanggal Ditemukan",
    "url": "Package URL",
}

# Which columns each dataset can offer, and what's pre-selected by default.
AVAILABLE_COLUMNS = {
    DATASET_LELANG: ["no", "nama_paket", "lpse", "jenis_pengadaan", "status", "hps", "tanggal_ditemukan", "url"],
    DATASET_BERANDA: ["no", "nama_paket", "lpse", "bagian", "kategori", "hps", "akhir_pendaftaran", "tanggal_ditemukan", "url"],
}
DEFAULT_COLUMNS = AVAILABLE_COLUMNS  # same set is sensible as the default today


class ExcelServiceError(Exception):
    """Same pattern as LPSEScraperError: a short Indonesian message safe to
    show directly to admin staff, plus technical detail for a collapsible
    "Detail Error" section."""

    def __init__(self, user_message: str, technical_detail: str = ""):
        super().__init__(user_message)
        self.user_message = user_message
        self.technical_detail = technical_detail or user_message


@dataclass
class ExcelExportResult:
    rows_written: int
    mode: str
    backup_path: str
    file_path: str
    sheet_name: str
    table_applied: bool
    table_skip_reason: Optional[str] = None


def _package_key(row) -> str:
    return row["package_id"] or row["fingerprint"]


def _column_value(dataset: str, field_key: str, row, no_value: int):
    if field_key == "no":
        return no_value
    if field_key == "url":
        return row["package_url"]
    if field_key == "tanggal_ditemukan":
        return (row["first_seen_at"] or "")[:10] or None
    if field_key == "hps":
        return row["hps_value"]
    if field_key == "nama_paket":
        return row["nama_paket"]
    if field_key == "lpse":
        return row["region_identifier"]

    if dataset == DATASET_LELANG:
        if field_key == "jenis_pengadaan":
            return row["jenis_pengadaan"]
        if field_key == "status":
            return row["tahapan"]
    else:  # DATASET_BERANDA
        if field_key == "bagian":
            return row["section"]
        if field_key == "kategori":
            return row["kategori"]
        if field_key == "akhir_pendaftaran":
            return row["akhir_pendaftaran_text"]

    return None


def backup_file(file_path: Path) -> Path:
    """Copy file_path into backups/ with a timestamped name, BEFORE any
    write. Never skipped - see module docstring rule 1."""
    BACKUPS_DIR.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().strftime("%Y-%m-%d_%H%M%S")
    backup_path = BACKUPS_DIR / f"{file_path.stem}_{timestamp}{file_path.suffix}"
    shutil.copy2(file_path, backup_path)
    return backup_path


def discover_excel_files(base_dir: Path = PROJECT_ROOT) -> list:
    """Find .xlsx files already sitting inside the project folder, so
    someone who just drops their workbook in there doesn't have to look
    up and type/paste a full path. Scoped to only this project folder -
    scanning the whole computer would be slow and would mean guessing at
    files the app has no business touching.

    Skips: the `backups/` folder (our own timestamped copies, not a
    workbook to export into), `venv/`, `.git/`, and Excel's own temporary
    lock files (start with "~$", created while a workbook is open).
    Returns paths sorted by most-recently-modified first, as plain
    strings (relative to base_dir when possible) - a good default guess
    is "whichever workbook was touched most recently".
    """
    skip_dirs = {"backups", "venv", ".venv", ".git", "__pycache__", "node_modules", "data"}
    found = []
    for path in base_dir.rglob("*.xlsx"):
        if path.name.startswith("~$"):
            continue
        if skip_dirs & set(path.relative_to(base_dir).parts[:-1]):
            continue
        found.append(path)
    found.sort(key=lambda p: p.stat().st_mtime, reverse=True)
    try:
        return [str(p.relative_to(base_dir)) for p in found]
    except ValueError:
        return [str(p) for p in found]


def create_new_excel_file(file_path: Path, sheet_name: str, start_cell: str, enabled_columns: list) -> None:
    """Convenience action for a user with no existing workbook yet: creates
    one with a header row directly above start_cell. This is the ONLY
    place this module ever writes a header row - the normal export path
    (export_packages) never touches anything above start_cell, since a
    real office template's headers must be preserved exactly as they are.
    """
    if file_path.exists():
        raise ExcelServiceError(
            f"File '{file_path.name}' sudah ada - tidak akan ditimpa. Gunakan file yang sudah ada, "
            "atau pilih nama/lokasi lain.",
            f"{file_path} already exists.",
        )
    file_path.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name

    start_row, start_col = coordinate_to_tuple(start_cell)
    header_row = start_row - 1
    if header_row >= 1:
        for col_offset, field_key in enumerate(enabled_columns):
            ws.cell(row=header_row, column=start_col + col_offset).value = COLUMN_LABELS.get(field_key, field_key)

    wb.save(file_path)


def _ranges_overlap(ref_a: str, ref_b: str) -> bool:
    a_min_col, a_min_row, a_max_col, a_max_row = range_boundaries(ref_a)
    b_min_col, b_min_row, b_max_col, b_max_row = range_boundaries(ref_b)
    return not (
        a_max_col < b_min_col or b_max_col < a_min_col
        or a_max_row < b_min_row or b_max_row < a_min_row
    )


def _apply_excel_table(ws, table_name: str, header_row: int, start_col: int, enabled_columns: list, last_data_row: int):
    """Write or refresh a native Excel Table (banded rows + filter
    dropdowns) over the header + data range, so an export hands back a
    ready-made table instead of raw values - and running this again later
    (with more, fewer, or the same rows) updates THIS SAME table's range
    rather than leaving a stale one behind or piling up "Table1",
    "Table2", ... duplicates.

    Returns (applied: bool, skip_reason: Optional[str]).
    """
    if last_data_row < header_row + 1:
        return False, "Belum ada baris data untuk dijadikan Tabel."

    end_col = start_col + len(enabled_columns) - 1
    ref = f"{get_column_letter(start_col)}{header_row}:{get_column_letter(end_col)}{last_data_row}"

    # Excel requires every header cell in the table to have real, non-blank
    # text. Use whatever is ACTUALLY in the header row (respecting a user's
    # own template wording) rather than forcing our own labels - but if any
    # of those cells are blank, we can't safely table-ify without writing
    # into a row this module promises never to touch.
    header_texts = []
    for col_offset in range(len(enabled_columns)):
        cell_value = ws.cell(row=header_row, column=start_col + col_offset).value
        if not cell_value:
            return False, (
                f"Baris header (baris {header_row}) punya sel kosong - Tabel Excel tidak dibuat agar baris "
                "header tidak diubah. Isi seluruh header terlebih dahulu, atau gunakan tombol 'Buat File Baru'."
            )
        header_texts.append(str(cell_value))

    # If a DIFFERENT table already covers overlapping cells (e.g. the user
    # turned this range into a Table by hand via Excel's own UI), leave it
    # alone rather than risk two overlapping tables corrupting the file.
    # Note: ws.tables.items() yields (name, ref_string) pairs, not Table
    # objects - the ref string is all _ranges_overlap needs anyway.
    for existing_name, existing_ref in list(ws.tables.items()):
        if existing_name == table_name:
            continue
        if _ranges_overlap(ref, existing_ref):
            return False, (
                f"Ada Tabel Excel lain ('{existing_name}') yang menempati sel yang sama - dibiarkan apa "
                "adanya. Data tetap ditulis sebagai nilai biasa."
            )

    if table_name in ws.tables:
        del ws.tables[table_name]

    columns = [TableColumn(id=i + 1, name=text) for i, text in enumerate(header_texts)]
    style = TableStyleInfo(
        name="TableStyleMedium7", showFirstColumn=False, showLastColumn=False,
        showRowStripes=True, showColumnStripes=False,
    )
    ws.add_table(Table(displayName=table_name, ref=ref, tableColumns=columns, tableStyleInfo=style))
    return True, None


def export_packages(conn: sqlite3.Connection, packages: list, config: dict, dataset: str = DATASET_LELANG) -> ExcelExportResult:
    """Write `packages` (already the relevant, earliest-first list for
    `dataset`) into the configured workbook/sheet/cell, in the given mode.
    Always backs up first. Raises ExcelServiceError with a ready-to-display
    Indonesian message on any problem (missing file, file locked, bad start
    cell, ...)."""

    file_path_str = config.get("file_path")
    if not file_path_str:
        raise ExcelServiceError("Belum ada file Excel yang dipilih. Atur di bagian atas tab Excel.")

    file_path = Path(file_path_str)
    if not file_path.exists():
        raise ExcelServiceError(
            f"File Excel tidak ditemukan: {file_path}. Periksa lokasinya, atau buat file baru dahulu.",
            f"Path does not exist: {file_path}",
        )

    sheet_name = config.get("sheet_name") or DEFAULT_SHEET_NAMES[dataset]
    start_cell = config.get("start_cell") or "A5"
    enabled_columns = config.get("enabled_columns") or DEFAULT_COLUMNS[dataset]
    mode = config.get("mode") or "replace"
    use_excel_table = config.get("use_excel_table", True)

    try:
        start_row, start_col = coordinate_to_tuple(start_cell)
    except Exception as exc:
        raise ExcelServiceError(
            f"Sel awal '{start_cell}' tidak valid. Gunakan format seperti 'A5'.", str(exc)
        ) from exc

    backup_path = backup_file(file_path)

    try:
        wb = openpyxl.load_workbook(file_path)
    except PermissionError as exc:
        raise ExcelServiceError(
            f"File Excel '{file_path.name}' sedang terbuka di program lain (mis. Microsoft Excel). "
            "Tutup file tersebut lalu coba lagi.",
            str(exc),
        ) from exc
    except Exception as exc:
        raise ExcelServiceError(
            f"Gagal membuka file Excel '{file_path.name}'. Pastikan ini file .xlsx yang valid.", str(exc)
        ) from exc

    ws = wb[sheet_name] if sheet_name in wb.sheetnames else wb.create_sheet(sheet_name)

    tracked_rows = db.get_tracked_excel_rows(conn, str(file_path), sheet_name)
    tracked_keys = {r["package_key"] for r in tracked_rows}

    if mode == "replace":
        for tr in tracked_rows:
            for col_offset in range(len(enabled_columns)):
                ws.cell(row=tr["row_number"], column=start_col + col_offset).value = None
        db.clear_tracked_excel_rows(conn, str(file_path), sheet_name)

        rows_to_write = list(packages)
        write_start_row = start_row
        no_start = 1
    elif mode == "append_new":
        rows_to_write = [p for p in packages if _package_key(p) not in tracked_keys]
        write_start_row = (max(tr["row_number"] for tr in tracked_rows) + 1) if tracked_rows else start_row
        no_start = len(tracked_rows) + 1
    else:
        raise ExcelServiceError(f"Mode Excel tidak dikenali: '{mode}'.")

    last_written_row = write_start_row - 1
    for i, pkg_row in enumerate(rows_to_write):
        current_row = write_start_row + i
        for col_offset, field_key in enumerate(enabled_columns):
            value = _column_value(dataset, field_key, pkg_row, no_start + i)
            ws.cell(row=current_row, column=start_col + col_offset).value = value
        db.add_tracked_excel_row(conn, str(file_path), sheet_name, _package_key(pkg_row), current_row)
        last_written_row = current_row

    # The table should always span every row THIS APP has ever written to
    # this sheet, not just the ones from this particular export - so use
    # the tracked-row bookkeeping (freshly updated above) as the source of
    # truth for where the data actually ends.
    all_tracked_rows = db.get_tracked_excel_rows(conn, str(file_path), sheet_name)
    last_data_row = max((r["row_number"] for r in all_tracked_rows), default=last_written_row)

    table_applied, table_skip_reason = (False, None)
    if use_excel_table:
        header_row = start_row - 1
        table_applied, table_skip_reason = _apply_excel_table(
            ws, TABLE_NAMES.get(dataset, "TabelLPSEMonitor"), header_row, start_col, enabled_columns, last_data_row,
        )

    try:
        wb.save(file_path)
    except PermissionError as exc:
        raise ExcelServiceError(
            f"File Excel '{file_path.name}' sedang terbuka di program lain (mis. Microsoft Excel). "
            "Tutup file tersebut lalu coba lagi. Data BELUM ditulis (backup sudah dibuat, file asli aman).",
            str(exc),
        ) from exc

    return ExcelExportResult(
        rows_written=len(rows_to_write), mode=mode, backup_path=str(backup_path),
        file_path=str(file_path), sheet_name=sheet_name,
        table_applied=table_applied, table_skip_reason=table_skip_reason,
    )
